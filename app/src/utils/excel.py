import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
import io

def read_urls_from_excel(uploaded_file):
    workbook = openpyxl.load_workbook(uploaded_file)
    sheet = workbook.worksheets[0] 
    urls = []
    
    for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0]:
            urls.append(str(row[0]).strip())
            
    return urls

def generate_excel_in_memory(data):
    wb = openpyxl.Workbook()
    ws = wb.worksheets[0] 
    ws.title = "Scrape Results"
    
    headers = ['KOL Type', 'KOL', 'Date', 'Week', 'Month', 'Platform', 'Link', 'Game', 'Views', 'Comments', 'Likes', 'Timestamp']
    ws.append(headers)
    
    header_fill = PatternFill(start_color="c5a8f0", end_color="8b52e0", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        
    for row in data:
        ws.append(row)
        
    for col_idx, col in enumerate(ws.columns, 1):
        max_length = 0
        column_letter = get_column_letter(col_idx)
        
        for cell in col:
            if cell.value is not None:
                try:
                    cell_length = len(str(cell.value))
                    if cell_length > max_length:
                        max_length = cell_length
                except:
                    pass
                    
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output